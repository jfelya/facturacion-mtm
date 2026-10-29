import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.pipeline import run_job, scan_sources
from app.writer import HEADERS, write_output
from tests.fixtures import make_claim, make_drivers, make_resumen, make_trips


class PipelineSmokeTests(unittest.TestCase):
    def test_one_day_batch_fills_and_exports(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            sources = root / "sources"
            sources.mkdir()

            claim_path = make_claim(
                root / "ClaimInformation test.xlsx",
                [{
                    "Claim Status": "Open",
                    "Trip Number": "M4F32EG2N0A",
                    "Appt. Date": "10/05/2026",
                    "Member Info": "LISA HALSEY",
                    "Driver": "",
                    "VIN": "",
                    "Provider Cost": "$15.00",
                    "Name (First and Last)": "LISA HALSEY",
                }],
            )
            make_resumen(
                sources / "RESUMEN FACTURACION 10-05-2026.xlsx",
                [{
                    "DATE": "10/05/2026",
                    "LOS": "AMB",
                    "TRIP ID": "",
                    "PAYER": "MTM",
                    "NAME": "LISA",
                    "LAST NAME": "HALSEY",
                    "DRIVER": 200,
                    "MILES": 5.67,
                    "NOTES": "",
                }],
            )
            make_trips(
                sources / "10-05.xlsx",
                [{
                    "Trip Number": "M4F32EG2N0A",
                    "Appointment Date": "10/05/2026",
                    "Time": "10:00 AM",
                    "Trip Mileage": 5.67,
                    "Trip Cost": 15.00,
                    "Member's First Name": "LISA",
                    "Member's Last Name": "HALSEY",
                    "Pickup City": "OCALA",
                }],
            )
            make_drivers(
                sources / "Informacion de Drivers.xlsx",
                [{
                    "DRIVER": "LISA DRIVER",
                    "COUNTY": "MARION",
                    "LICENSE NUMBER": "L1",
                    "VIN Number": "V1",
                    "CURRENT VEHICLE": "",
                    "VEHICLE PLATE": "",
                    "NUMERO": 200,
                }],
            )

            res = run_job(claim_path, sources)
            self.assertNotIn("error", res, res.get("error"))
            row = res["claim"]["rows"][0]
            self.assertEqual(row["driver"], "200")
            self.assertEqual(row["license"], "L1")
            self.assertEqual(row["vin"], "V1")
            self.assertEqual(row["milage"], 5.67)
            self.assertIsNotNone(row["pu"])
            self.assertIsNotNone(row["dropoff"])
            self.assertLess(row["pu"], row["dropoff"])

            drivers_file = sources / "Informacion de Drivers.xlsx"
            src = scan_sources(sources, drivers_file)
            res2 = run_job(claim_path, sources, drivers_path=drivers_file, sources=src)
            self.assertEqual(res["report"]["stats"], res2["report"]["stats"])
            self.assertEqual([r["trip"] for r in res["claim"]["rows"]],
                             [r["trip"] for r in res2["claim"]["rows"]])

            out = root / "out.xlsx"
            info = write_output(res["claim"], out)
            self.assertEqual(info["rows"], 1)
            self.assertEqual(info["sheets"], 1)

            wb = load_workbook(out, read_only=True)
            try:
                self.assertIn("10-05", wb.sheetnames)
                values = list(wb["10-05"].iter_rows(values_only=True))
            finally:
                wb.close()

            self.assertEqual(tuple(values[0]), tuple(HEADERS))
            self.assertEqual(len(values), 2)
            self.assertEqual(values[1][0], 200)
            self.assertEqual(values[1][3], "10/05/2026")
            self.assertEqual(values[1][4], "M4F32EG2N0A")
            self.assertEqual(values[1][5], "LISA HALSEY")
            self.assertEqual(values[1][8], 5.67)
            self.assertEqual(values[1][9], "$15.00")


if __name__ == "__main__":
    unittest.main()
