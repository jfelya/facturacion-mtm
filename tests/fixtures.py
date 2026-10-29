"""Synthetic fixtures for the test suite. No production data is read."""
from pathlib import Path

from docx import Document
from openpyxl import Workbook

CLAIM_HEADERS = [
    "Claim Status", "Trip Number", "Appt. Date", "Member Info",
    "Driver", "VIN", "Provider Cost", "Name (First and Last)",
]
RESUMEN_HEADERS = [
    "DATE", "LOS", "TRIP ID", "PAYER", "NAME", "LAST NAME",
    "DRIVER", "MILES", "NOTES",
]
TRIPS_HEADERS = [
    "Trip Number", "Appointment Date", "Time", "Trip Mileage",
    "Trip Cost", "Member's First Name", "Member's Last Name", "Pickup City",
]
DRIVERS_HEADERS = [
    "DRIVER", "COUNTY", "LICENSE NUMBER", "VIN Number",
    "CURRENT VEHICLE", "VEHICLE PLATE", "NUMERO",
]


def _write_xlsx(path, headers, rows, sheet="Sheet1"):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        if isinstance(row, dict):
            ws.append([row.get(h, "") for h in headers])
        else:
            ws.append(list(row))
    wb.save(path)
    return path


def make_claim(path, rows):
    return _write_xlsx(path, CLAIM_HEADERS, rows)


def make_resumen(path, rows):
    return _write_xlsx(path, RESUMEN_HEADERS, rows)


def make_trips(path, rows):
    return _write_xlsx(path, TRIPS_HEADERS, rows)


def make_drivers(path, rows):
    return _write_xlsx(path, DRIVERS_HEADERS, rows, sheet="Vehiculos")


def make_docx_template(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    table = doc.add_table(rows=15, cols=8)
    table.cell(0, 0).text = "Date of Service"
    table.cell(0, 1).text = "Trip number"
    table.cell(0, 2).text = "Recipient's Name"
    table.cell(0, 3).text = "PickUp Time"
    table.cell(0, 4).text = "DropOff Time"
    table.cell(0, 5).text = "Milleage"
    table.cell(0, 6).text = "Per Trip Bill Amount"
    table.cell(0, 7).text = "Recipient's Signature"
    doc.save(path)
    return path
