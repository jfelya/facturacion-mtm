import tempfile
import unittest
from datetime import time as dtime
from pathlib import Path

from openpyxl import load_workbook

from app.bulk import generate_bulks, template_path
from app.driverlogs import generate_driverlogs
from app.report import report_text
from app.ui import status_of
from app.writer import write_output
from tests.fixtures import make_docx_template


def _rows():
    return [
        {
            "row": 2, "date": "2026-10-05", "date_text": "10/05/2026",
            "trip": "N0A", "leg": "A", "member": "NORMAL MEMBER",
            "member_raw": "NORMAL MEMBER", "driver": "200", "license": "L200",
            "vin": "V200", "pu": dtime(9, 0), "dropoff": dtime(9, 20),
            "pu_appt": dtime(9, 30), "milage": 5.0, "cost": "$10.00", "flag": "",
        },
        {
            "row": 3, "date": "2026-10-05", "date_text": "10/05/2026",
            "trip": "B0A", "leg": "A", "member": "BILLED MEMBER",
            "member_raw": "BILLED MEMBER", "driver": "220", "license": "L220",
            "vin": "V220", "pu": dtime(10, 0), "dropoff": dtime(10, 20),
            "pu_appt": dtime(10, 30), "milage": 6.0, "cost": "$11.00",
            "flag": "facturado",
        },
    ]


def _report():
    return {
        "stats": {
            "rows": 2, "id_matched": 0, "order_assigned": 1, "no_data": 0,
            "flagged_unapproved": 0, "note_overrides": 0, "pu_filled": 1,
            "pu_willcall": 0, "pu_missing": 0, "already_billed": 1,
        },
        "warnings": [],
        "review": [],
        "already_billed": [{"date": "2026-10-05", "trip": "B0A",
                            "member": "BILLED MEMBER"}],
        "needs_driver": [],
        "driver_candidates": {},
        "suggested": [],
        "not_billed": [],
        "resumen_without_claim": [],
    }


class FacturadoOutputTests(unittest.TestCase):
    def test_writer_excludes_facturado(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.xlsx"
            info = write_output({"rows": _rows()}, out)
            self.assertEqual(info["skipped"], 1)
            self.assertEqual(info["rows"], 1)
            wb = load_workbook(out, read_only=True)
            try:
                values = list(wb["10-05"].iter_rows(values_only=True))
            finally:
                wb.close()
        self.assertEqual(len(values), 2)
        self.assertEqual(values[1][0], 200)

    def test_bulk_excludes_facturado_and_keeps_template(self):
        before = template_path().read_bytes()
        with tempfile.TemporaryDirectory() as tmp:
            res = generate_bulks(_rows(), tmp)
        self.assertEqual(len(res["created"]), 1)
        self.assertEqual(len(res["excluded"]), 1)
        self.assertEqual(template_path().read_bytes(), before)

    def test_driverlogs_exclude_facturado(self):
        with tempfile.TemporaryDirectory() as tmp:
            templates = Path(tmp) / "templates"
            make_docx_template(templates / "200 NORMAL DRIVERLOG VACIO.docx")
            make_docx_template(templates / "220 BILLED DRIVERLOG VACIO.docx")
            res = generate_driverlogs(_rows(), templates, Path(tmp) / "days")
            self.assertEqual(res["billed_skipped"], 1)
            self.assertTrue(all("220" not in path for path in res["created"]))
            self.assertEqual(len(res["created"]), 1)

    def test_report_lists_billed(self):
        text = report_text({"report": _report(), "claim_path": "x.xlsx"}, "es")
        self.assertIn("YA FACTURADO", text)

    def test_status_is_facturado(self):
        row = {"driver": "220", "milage": 1.0, "flag": "facturado"}
        self.assertEqual(status_of(row), "FACTURADO")


if __name__ == "__main__":
    unittest.main()
