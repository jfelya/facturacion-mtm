"""Fill the Word driverlogs (.docx) from the final claim.

One driverlog **per driver per day**, copied from the templates folder and written to
`<days_root>\\<day>\\Docs\\<day> <NUM>.docx` (e.g. `10-05 200.docx`).

Rules (confirmed by the user, NOTES §15):
- The templates (`<NUM> <NOMBRE> DRIVERLOG VACIO.docx`) are **never modified** — a filled copy is written.
- Only the table's data rows are written; everything else stays untouched.
- Columns: Date of Service | Trip number | Recipient's Name | PickUp Time | DropOff Time |
  Milleage | Per Trip Bill Amount | Recipient's Signature (blank).
- Drivers without a template are skipped (they were already replaced in the claim, NOTES §6.5).
"""
import re
from pathlib import Path

_HEADER_COLS = 8   # first 8 cells carry the data (there is a trailing empty cell)
MAX_TRIPS = 14     # data rows in the template (only grows if a driver ever exceeds it)


def scan_templates(folder):
    """{driver_number: template_path} from files named `<NUM> <NAME> ... .docx`."""
    out = {}
    p = Path(folder) if folder else None
    if not p or not p.is_dir():
        return out
    for f in sorted(p.glob("*.docx")):
        if f.name.startswith("~$"):
            continue
        m = re.match(r"\s*(\d+)\b", f.stem)
        if m:
            out.setdefault(m.group(1), f)
    return out


def day_name(iso):
    """'2026-10-05' -> '10-05'."""
    try:
        _y, m, d = str(iso).split("-")
        return f"{m}-{d}"
    except Exception:
        return str(iso)


def _hhmm(v):
    if v is None or v == "":
        return ""
    if hasattr(v, "hour"):                 # datetime.time
        return f"{v.hour}:{v.minute:02d}"
    m = re.match(r"(\d{1,2}):(\d{2})", str(v).strip())
    return f"{int(m.group(1))}:{m.group(2)}" if m else str(v).strip()


def _num(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _row_values(r):
    return [
        r.get("date_text") or "",                       # Date of Service
        r.get("trip") or "",                            # Trip number
        r.get("member_raw") or r.get("member") or "",   # Recipient's Name
        _hhmm(r.get("pu")),                             # PickUp Time
        _hhmm(r.get("dropoff")),                        # DropOff Time
        _num(r.get("milage")),                          # Milleage
        r.get("cost") or "",                            # Per Trip Bill Amount
        "",                                             # Recipient's Signature
    ]


def fill_doc(template_path, trips):
    """Open the template, wipe the data rows and write `trips`. Returns a Document."""
    from docx import Document

    doc = Document(str(template_path))
    if not doc.tables:
        raise ValueError("la plantilla no tiene tabla")
    table = doc.tables[0]

    # wipe every data row (row 0 is the header)
    for row in table.rows[1:]:
        for cell in row.cells[:_HEADER_COLS]:
            cell.text = ""

    # make sure there are enough data rows
    while len(table.rows) - 1 < len(trips):
        table.add_row()

    # write
    for i, r in enumerate(trips):
        cells = table.rows[1 + i].cells
        for cell, val in zip(cells[:_HEADER_COLS], _row_values(r)):
            if cell.text != val:
                cell.text = val
    return doc


def generate_driverlogs(rows, templates_folder, days_root, only_days=None):
    """Write one filled driverlog per driver per day.

    rows            claim rows (dicts with date, driver, trip, date_text, member, pu, dropoff, milage, cost)
    templates_folder folder with `<NUM> ... DRIVERLOG VACIO.docx`
    days_root       root that holds the day folders (e.g. ...\\1. PARA FACTURAR)
    only_days       optional set of day names ('10-05') to limit the run

    Returns {'created': [...], 'skipped': [(day, drv)], 'errors': [...], 'days': {...}}
    """
    tpl = scan_templates(templates_folder)
    if not tpl:
        return {"created": [], "skipped": [], "errors": [f"No hay plantillas .docx en {templates_folder}"],
                "days": {}, "templates": 0}
    if not days_root:
        return {"created": [], "skipped": [], "errors": ["Falta la carpeta de días (days_root)"],
                "days": {}, "templates": len(tpl)}

    by_day = {}
    for r in rows:
        d = r.get("date") or ""
        if d:
            by_day.setdefault(day_name(d), []).append(r)

    created, skipped, errors, days = [], [], [], {}
    for day in sorted(by_day):
        if only_days and day not in only_days:
            continue
        by_drv = {}
        for r in by_day[day]:
            drv = str(r.get("driver") or "").strip()
            if drv:
                by_drv.setdefault(drv, []).append(r)
        docs_dir = Path(days_root) / day / "Docs"
        made = 0
        for drv in sorted(by_drv, key=lambda x: (len(x), x)):
            template = tpl.get(drv)
            if not template:
                skipped.append((day, drv))
                continue
            try:
                docs_dir.mkdir(parents=True, exist_ok=True)
                doc = fill_doc(template, by_drv[drv])
                out = docs_dir / f"{day} {drv}.docx"
                doc.save(str(out))
                created.append(str(out))
                made += 1
            except Exception as e:                       # noqa: BLE001
                errors.append(f"{day} {drv}: {type(e).__name__}: {e}")
        days[day] = {"docs": made, "drivers": sorted(by_drv, key=lambda x: (len(x), x))}

    return {"created": created, "skipped": skipped, "errors": errors, "days": days,
            "templates": len(tpl)}


def report_lines(res, lang="es"):
    """Small text summary for the UI/report."""
    es = lang == "es"
    L = []
    L.append(("DRIVERLOGS" if es else "DRIVERLOGS"))
    L.append((f"  Plantillas encontradas: {res.get('templates', 0)}"
              if es else f"  Templates found: {res.get('templates', 0)}"))
    L.append((f"  Documentos generados: {len(res.get('created', []))}"
              if es else f"  Documents created: {len(res.get('created', []))}"))
    for d in sorted(res.get("days", {})):
        info = res["days"][d]
        L.append(f"    {d}: {info['docs']} doc(s) -> drivers {', '.join(info['drivers'])}")
    if res.get("skipped"):
        L.append(("  Sin plantilla (omitidos): " if es else "  No template (skipped): ")
                 + ", ".join(f"{d} {n}" for d, n in res["skipped"]))
    for e in res.get("errors", []):
        L.append(("  ERROR: " if es else "  ERROR: ") + e)
    return "\n".join(L)
