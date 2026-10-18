"""Generate the MTM **BULK UPLOAD** Excel — one file per day — from the final claim.

Sheet `Export Trip`, 10 columns (exact template order):

    TripNumber | Driver | Vehicle | ScheduledPickupTime | ReportedPickupArriveTime |
    ReportedPickupPerformTime | ScheduledDropoffTime | ReportedDropoffArriveTime |
    ReportedDropoffPerformTime | TripLogSignature

Meanings (per the user's `BULK UPLOAD Training.xlsx`):

    Driver               = the driver's **License**  (claim "License")
    Vehicle              = the vehicle **VIN**        (claim "VIN")
    ScheduledPickupTime  = planned pickup             (claim "PU Time")
    ScheduledDropoffTime = the **trip file's `Time`** (= the appointment, 15-min
                           marks). For a **Will Call** (no appointment) it is the
                           computed return dropoff (claim `DropOff Time`).
    TripLogSignature     = "Y" (always)

The four **Reported** columns do NOT exist in the ClaimInformation, so they are
**simulated** with a deterministic jitter (same trip -> same times on every run,
so the output is idempotent), matching the archived examples:

    ReportedPickupArriveTime   = PU  - 1..5 min   (A leg) / PU - 1..3 min (B leg)
    ReportedPickupPerformTime  = arrive + 1 min   (patient gets in)
    ReportedDropoffArriveTime  = SchedDrop - 1..13 min
    ReportedDropoffPerformTime = arrive (same)    (patient gets out)

Extra safety rules (user, 18-oct):

* A **B leg** (return) never arrives for its pickup before the matching **A leg**
  has dropped the member off — otherwise a back-to-back rider would "teleport".
* The drop-off can never happen before the pick-up on the same trip.
* B legs only arrive 1..3 min early (not more) precisely to avoid that clash.

Output goes to `<days_root>\\<MM-DD>\\BULK UPLOAD <MM-DD-YYYY>.xlsx`, formatted
from the bundled template (`app/assets/bulk_template.xlsx`) so the layout stays
identical. The template file is NEVER modified.
"""
import random
import sys
import zlib
from datetime import date, datetime, time as dtime, timedelta
from pathlib import Path

SHEET = "Export Trip"
HEADERS = [
    "TripNumber", "Driver", "Vehicle", "ScheduledPickupTime",
    "ReportedPickupArriveTime", "ReportedPickupPerformTime",
    "ScheduledDropoffTime", "ReportedDropoffArriveTime",
    "ReportedDropoffPerformTime", "TripLogSignature",
]
SIGNATURE = "Y"
FIRST_DATA_ROW = 2


# --------------------------------------------------------------------------- helpers
def _asset_path(name):
    base = getattr(sys, "_MEIPASS", None)
    if base:
        p = Path(base) / "app" / "assets" / name
        if p.exists():
            return p
    return Path(__file__).resolve().parent / "assets" / name


def template_path():
    """The bundled empty BULK UPLOAD template used as the format master."""
    return _asset_path("bulk_template.xlsx")


def day_name(iso):
    """'2026-10-05' -> '10-05'."""
    try:
        _y, m, d = str(iso).split("-")
        return f"{m}-{d}"
    except Exception:
        return str(iso)


def bulk_filename(iso):
    """'2026-10-05' -> 'BULK UPLOAD 10-05-2026.xlsx'."""
    try:
        y, m, d = str(iso).split("-")
        return f"BULK UPLOAD {m}-{d}-{y}.xlsx"
    except Exception:
        return "BULK UPLOAD.xlsx"


def _as_time(v):
    """Accept datetime, time or an 'H:MM'/'HH:MM' string; return time or None."""
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, dtime):
        return v
    if isinstance(v, str) and ":" in v:
        try:
            h, m = v.strip().split(":")[:2]
            return dtime(int(h), int(m[:2]))
        except Exception:
            return None
    return None


def _add(t, minutes):
    return (datetime.combine(date(2000, 1, 1), t) + timedelta(minutes=minutes)).time()


def _mins(t):
    return t.hour * 60 + t.minute


def scheduled_dropoff(row):
    """Scheduled drop-off for the BULK.

    * Normal trip: the trip file's `Time` (the appointment) -> `pu_appt`
      (a 15-minute mark; it is also what the pickup time is derived from).
    * Will Call: the trip file has `Time = "Will Call"`, so the app derives
      `A + 2h`. That is the **return pickup**, not an appointment -> use the
      computed return dropoff (the claim's `DropOff Time`) instead, so the
      trip stays physically consistent.

    Detection: for a Will Call `pu_appt` sits only 1..3 min after `pu`; for a
    normal trip it is `pu + (lead + margin)` (>= ~12 min).
    """
    appt = _as_time(row.get("pu_appt"))
    drop = _as_time(row.get("dropoff"))
    pu = _as_time(row.get("pu"))
    if appt and drop and pu:
        gap = _mins(appt) - _mins(pu)
        if 0 <= gap <= 5:                 # Will Call: A+2h is the return pickup
            return drop
    return appt or drop


def _leg(row):
    leg = (row.get("leg") or "").upper()
    if leg in ("A", "B"):
        return leg
    trip = str(row.get("trip") or "").upper()
    return trip[-1] if trip[-1:] in ("A", "B") else ""


# --------------------------------------------------------------------------- simulation
def _base_sim(row):
    """Base simulated (arrive_pu, perform_pu, arrive_drop, perform_drop) for one trip."""
    pu = _as_time(row.get("pu"))
    drop = scheduled_dropoff(row)
    seed = zlib.crc32(f"{row.get('date', '')}|{row.get('trip', '')}|bulk".encode("utf-8"))
    rng = random.Random(seed)
    # B legs (returns) arrive only 1..3 min early so they never undercut the A leg.
    lo, hi = (1, 3) if _leg(row) == "B" else (1, 5)

    arrive_pu = perform_pu = arrive_drop = perform_drop = None
    if pu:
        pu = pu.replace(second=0, microsecond=0)
        arrive_pu = _add(pu, -rng.randint(lo, hi))
        perform_pu = _add(arrive_pu, 1)
    if drop:
        drop = drop.replace(second=0, microsecond=0)
        arrive_drop = _add(drop, -rng.randint(1, 13))
        perform_drop = arrive_drop
    # the patient cannot get out before he/she got in
    if perform_pu and arrive_drop and _mins(arrive_drop) <= _mins(perform_pu):
        arrive_drop = _add(perform_pu, rng.randint(4, 12))
        perform_drop = arrive_drop
    return [arrive_pu, perform_pu, arrive_drop, perform_drop]


def sim_times(row):
    """Single-trip simulated times (no A/B context). Tuple for convenience."""
    return tuple(_base_sim(row))


def day_sims(rows):
    """Simulate a whole day, keeping B legs from arriving before their A leg drops off.

    Returns a list aligned with `rows`; each item is [arrive_pu, perform_pu, arrive_drop, perform_drop].
    """
    sims = [_base_sim(r) for r in rows]
    a_by_key = {}
    for i, r in enumerate(rows):
        trip = str(r.get("trip") or "").upper()
        if trip.endswith("A"):
            a_by_key[(r.get("date"), trip[:-1])] = sims[i]
    for i, r in enumerate(rows):
        trip = str(r.get("trip") or "").upper()
        if not trip.endswith("B"):
            continue
        a = a_by_key.get((r.get("date"), trip[:-1]))
        if not a:
            continue
        a_out = a[3] or a[2]                      # A leg: member got out
        if a_out and sims[i][0] and _mins(sims[i][0]) < _mins(a_out):
            sims[i][0] = a_out                    # B driver only arrives after dropping off
            sims[i][1] = _add(a_out, 1)
            if sims[i][3] and _mins(sims[i][3]) < _mins(sims[i][1]):
                sims[i][3] = _add(sims[i][1], 4)
                sims[i][2] = sims[i][3]
    return sims


def row_values(row, sim=None):
    """The 10 cell values for one claim row (A..J)."""
    a_pu, p_pu, a_dr, p_dr = sim if sim else _base_sim(row)
    return [
        row.get("trip") or "",              # A TripNumber
        row.get("license") or "",           # B Driver (= License)
        row.get("vin") or "",               # C Vehicle (= VIN)
        _as_time(row.get("pu")),            # D ScheduledPickupTime
        a_pu,                               # E ReportedPickupArriveTime
        p_pu,                               # F ReportedPickupPerformTime
        scheduled_dropoff(row),             # G ScheduledDropoffTime (trip-file Time)
        a_dr,                               # H ReportedDropoffArriveTime
        p_dr,                               # I ReportedDropoffPerformTime
        SIGNATURE,                          # J TripLogSignature
    ]


# --------------------------------------------------------------------------- workbook
def fill_workbook(rows, template=None):
    """Load the BULK UPLOAD template and write `rows` into it. Returns an openpyxl Workbook."""
    from copy import copy as _copy

    from openpyxl import load_workbook

    tpl = Path(template) if template else template_path()
    if not Path(tpl).exists():
        raise FileNotFoundError(f"No encuentro la plantilla BULK UPLOAD: {tpl}")
    wb = load_workbook(str(tpl))
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]

    # grow (cloning row 2's style) if a day ever has more trips than the template rows
    need_last = FIRST_DATA_ROW + len(rows) - 1
    while ws.max_row < max(need_last, FIRST_DATA_ROW):
        r = ws.max_row + 1
        for c in range(1, len(HEADERS) + 1):
            ws.cell(r, c)._style = _copy(ws.cell(FIRST_DATA_ROW, c)._style)

    # wipe every data cell (values + any leftover formula)
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        for c in range(1, len(HEADERS) + 1):
            ws.cell(r, c).value = None

    # write the rows (with the A->B clash handling)
    sims = day_sims(rows)
    for i, row in enumerate(rows):
        r = FIRST_DATA_ROW + i
        for c, val in enumerate(row_values(row, sims[i]), start=1):
            ws.cell(r, c).value = val

    # delete the leftover blank formatted rows ("borrar espacios en blanco")
    leftover = ws.max_row - (FIRST_DATA_ROW + len(rows) - 1)
    if leftover > 0:
        ws.delete_rows(FIRST_DATA_ROW + len(rows), leftover)
    return wb


def generate_bulks(rows, days_root, template=None, only_days=None):
    """Write one `BULK UPLOAD <date>.xlsx` per day into `<days_root>\\<MM-DD>\\`.

    Only billable rows (with a driver, not `nofacturar`) are included.

    Returns {'created': [...], 'days': {...}, 'excluded': [...], 'errors': [...], 'template': path}
    """
    selected = [r for r in rows
                if r.get("driver") and r.get("flag") != "nofacturar" and r.get("date")]
    excluded = [r for r in rows if r.get("date") and r not in selected]

    by_day = {}
    for r in selected:
        by_day.setdefault(r["date"], []).append(r)

    created, errors, days = [], [], {}
    if not days_root:
        return {"created": [], "days": {}, "excluded": excluded,
                "errors": ["Falta la carpeta de días (days_root)"], "template": str(template or template_path())}

    for iso in sorted(by_day):
        day = day_name(iso)
        if only_days and day not in only_days:
            continue
        day_rows = by_day[iso]
        try:
            wb = fill_workbook(day_rows, template)
            folder = Path(days_root) / day
            folder.mkdir(parents=True, exist_ok=True)
            out = folder / bulk_filename(iso)
            wb.save(str(out))
            created.append(str(out))
            days[day] = {"file": str(out), "rows": len(day_rows)}
        except Exception as e:                      # noqa: BLE001
            errors.append(f"{day}: {type(e).__name__}: {e}")

    return {"created": created, "days": days, "excluded": excluded, "errors": errors,
            "template": str(template or template_path())}


def report_lines(res, lang="es"):
    """Small text summary for the UI/report."""
    es = lang == "es"
    L = []
    L.append("BULK UPLOAD")
    L.append((f"  Archivos generados: {len(res.get('created', []))}"
              if es else f"  Files created: {len(res.get('created', []))}"))
    for d in sorted(res.get("days", {})):
        info = res["days"][d]
        L.append(f"    {d}: {info['rows']} fila(s) -> {Path(info['file']).name}")
    if res.get("excluded"):
        L.append((f"  Filas no incluidas (sin driver / no facturar): {len(res['excluded'])}"
                  if es else f"  Rows not included (no driver / not billed): {len(res['excluded'])}"))
    for e in res.get("errors", []):
        L.append("  ERROR: " + e)
    return "\n".join(L)
