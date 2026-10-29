"""Shared helpers: normalization and safe xlsx reading."""
import re
from datetime import datetime, date, time as dtime, timedelta
from openpyxl import load_workbook


def s(v):
    """Stringify cell value sanely (int-like floats -> '200')."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


_BILLED_NOTE = re.compile(r"^\s*(?:ya\s+)?facturado\b", re.I)
_ACTIONABLE_NOTE = re.compile(r"auto\s*\d+|duplicad|perfil|vehiculo|will\s*call", re.I)


def note_kind(note):
    """Classify a summary note: 'billed', 'actionable', 'info' or ''."""
    t = s(note)
    if not t:
        return ""
    if _BILLED_NOTE.match(t):
        return "billed"
    if _ACTIONABLE_NOTE.search(t):
        return "actionable"
    return "info"


def nonempty(rows):
    return [r for r in rows if any(v not in (None, "") for v in r)]


def read_rows(path, sheet=None):
    """Return (rows, sheet_title, sheetnames). Read-only, data-only."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        return rows, ws.title, list(wb.sheetnames)
    finally:
        wb.close()


def read_header(path):
    """Return the first row and sheet names only; cheaper than read_rows."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return list(row), list(wb.sheetnames)
    finally:
        wb.close()


def day_name(iso):
    """'2026-10-05' -> '10-05'; invalid input keeps the writer's old 'Hoja1'."""
    try:
        _y, m, d = str(iso).split("-")
        return f"{m}-{d}"
    except Exception:
        return "Hoja1"


def add_minutes(t, minutes):
    """Add minutes to a datetime.time, wrapping across midnight."""
    return (datetime.combine(date(2000, 1, 1), t) + timedelta(minutes=minutes)).time()


def numeric_key(value):
    """Sort numeric driver numbers first, then text."""
    t = str(value)
    return (0, int(t)) if t.isdigit() else (1, t)


def norm_name(*parts):
    txt = " ".join(s(p) for p in parts)
    txt = txt.upper().replace(".", "")
    return " ".join(txt.split())


def norm_num(v):
    """Normalize a driver number: 200/200.0/'200-1' -> '200'."""
    if isinstance(v, bool):
        return s(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int):
        v = str(v)
    t = s(v)
    m = re.fullmatch(r"(\d+)(?:-\d+)?", t)
    if m:
        return m.group(1)
    return t


def to_iso(v):
    """Any date-ish value -> 'YYYY-MM-DD' or None."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    t = s(v)
    if "/" in t:
        parts = t.split("/")
        if len(parts) == 3:
            mm, dd, yy = (p.strip() for p in parts)
            try:
                return f"{int(yy):04d}-{int(mm):02d}-{int(dd):02d}"
            except ValueError:
                return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", t)
    if m:
        return m.group(0)
    return None


def fmt_mdy(v):
    """Date-ish value -> 'MM/DD/YYYY' text (for the output column)."""
    if isinstance(v, datetime) or isinstance(v, date):
        return f"{v.month:02d}/{v.day:02d}/{v.year}"
    return s(v)


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_time_value(v):
    """Return ('time', datetime.time) | ('willcall', None) | (None, None)."""
    if isinstance(v, datetime):
        return ("time", v.time())
    if isinstance(v, dtime):
        return ("time", v)
    t = s(v)
    if not t:
        return (None, None)
    if "will call" in t.lower():
        return ("willcall", None)
    for fmt in ("%I:%M %p", "%I:%M:%S %p", "%H:%M", "%H:%M:%S", "%I:%M%p"):
        try:
            return ("time", datetime.strptime(t.upper(), fmt).time())
        except ValueError:
            pass
    return (None, None)
